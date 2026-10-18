"""The importable kinds (§51) — learners and term grades.

Each spec supplies its columns, a validator and a committer; the sequence
around them lives in `app/import_pipeline.py`.

Both validators load their reference data **once** for the whole file.
The database is ~85ms away, so checking each row with its own query would
make a 300-row masterlist take half a minute.

§51 names the errors these have to catch: duplicate LRN, unknown section,
unknown subject, invalid grade, impossible date, and a subject not
offered during that term. Each has a test.
"""

import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

import openpyxl

from app import audit_service
from app.display_time import SCHOOL_TZ
from app.grading_service import recompute_enrollment_grades_batch
from app.guardian_contact import (
    CONSENT_NEEDS_CONTACT,
    apply_contact,
    clean_email,
    consent_without_contact,
    contact_values,
    normalise_mobile,
    parse_consent,
)
from app.import_pipeline import (
    ColumnSpec,
    ImportSpec,
    RowError,
    ValidationResult,
    detect_date_order,
    parse_date,
    parse_grade,
    parse_lrn,
)
from app.learner_access import grade_12_learner_ids, may_edit
from app.models.academic_structure import Section
from app.models.enums import EnrollmentStatus, GradeWorkflowStatus, ImportJobType, Sex
from app.models.grades import TermGrade
from app.models.learners import Enrollment, Learner, LearnerAdmissionRecord
from app.models.organization import Term
from app.models.subjects import SectionSubjectOffering, Subject
from app.naming import normalize_name
from app.section_access import is_advised_by
from app.text_rules import ADDRESS_MAX, NOTE_MAX, SHORT_MAX

# `term_grades.source` is free text. Stamping it lets a migrated grade be
# told apart from one a teacher typed, which matters when reconciling the
# first year's data against the old workbook.
GRADE_SOURCE_IMPORT = "IMPORT"

# --- Learners --------------------------------------------------------------

# Parent/guardian contact (§78.1), shared by the learner import and the
# contact update below. All optional.
#
# Every alias names a parent or guardian. A bare "Email" or "Mobile" header
# is deliberately *not* an alias: a masterlist often carries the learner's
# own address under it, and auto-mapping that here would send the
# learner's term card and concern notices to the learner. Left unmatched,
# the person importing has to choose the column, which is the point.
GUARDIAN_COLUMNS = [
    ColumnSpec(
        "guardian_name", "Parent/Guardian Name", False,
        ("parentname", "guardianname", "parentguardian", "nameofparentorguardian"),
    ),
    ColumnSpec(
        "guardian_email", "Parent/Guardian Email", False,
        ("parentemail", "guardianemail", "parentemailaddress", "guardianemailaddress"),
    ),
    # A number Excel has turned into 9171234567 (leading zero dropped) is
    # repaired by `normalise_mobile` rather than refused.
    ColumnSpec(
        "guardian_mobile", "Parent/Guardian Mobile", False,
        (
            "parentmobile", "guardianmobile", "parentcontact", "guardiancontact",
            "parentcontactnumber", "guardiancontactnumber", "parentguardiancontactnumber",
            "contactnumberofparentorguardian", "contactnumberofparentguardian",
        ),
    ),
    # Yes/No. Blank is "not stated".
    ColumnSpec("notices_consent", "Notice Consent", False, ("parentconsent", "noticeconsent")),
]

LEARNER_COLUMNS = [
    ColumnSpec("last_name", "Last Name", True, ("surname", "familyname")),
    ColumnSpec("first_name", "First Name", True, ("givenname",)),
    ColumnSpec("middle_name", "Middle Name", False, ("middlename",)),
    ColumnSpec("extension_name", "Extension Name", False, ("suffix", "extname")),
    ColumnSpec("sex", "Sex", True, ("gender",)),
    ColumnSpec("birthdate", "Birthdate", True, ("dateofbirth", "dob", "birthday")),
    ColumnSpec("lrn", "LRN", False, ("learnerreferencenumber",)),
    # Optional. Filled in, the learner is enrolled into that section in
    # one step; left blank, they are created and enrolled later on the
    # Enrollment page, which is what this import did before.
    ColumnSpec("section", "Section", False, ("sectionname", "class")),
    *GUARDIAN_COLUMNS,
]


def _parse_sex(raw: str):
    value = str(raw or "").strip().upper()
    if value in ("M", "MALE"):
        return Sex.MALE, None
    if value in ("F", "FEMALE"):
        return Sex.FEMALE, None
    return None, f"{raw!r} is not MALE or FEMALE"


def _section_lookup(session, school_year_id, adviser_user_id=None) -> tuple[dict, set]:
    """Sections for the year, keyed by upper-cased name.

    A name is only unique per grade level, so the same name may legally
    exist in both Grade 11 and Grade 12. Those names are collected
    separately and rejected rather than guessed at — silently enrolling a
    Grade 11 learner into the Grade 12 section of the same name would be
    very hard to notice afterwards.

    `adviser_user_id` settles that tie-break rather than filtering the
    map: an adviser who holds only one of two same-named sections has
    named it unambiguously, since the other one is refused to them
    anyway. Sections they don't advise stay in the map **on purpose** —
    the caller can then say "that isn't your section" instead of the
    misleading "unknown section", which would send a teacher off to check
    a spelling that was right all along.
    """
    if school_year_id is None:
        return {}, set()

    candidates: dict[str, list] = {}
    for section in session.query(Section).filter_by(school_year_id=school_year_id).all():
        key = (section.name or "").strip().upper()
        candidates.setdefault(key, []).append(section)

    by_name: dict[str, object] = {}
    ambiguous: set[str] = set()
    for key, found in candidates.items():
        if adviser_user_id is not None:
            found = [s for s in found if is_advised_by(s, adviser_user_id)] or found
        if len(found) > 1:
            ambiguous.add(key)
        by_name[key] = found[0]
    return by_name, ambiguous


def _flag_long_names(result, number, fields) -> None:
    """Adds an error for each name longer than the data-entry limit. The
    edit forms cap names at SHORT_MAX; without the same check here a file
    could write what no form would accept, to a free-tier database."""
    for label, value in fields:
        if value and len(value) > SHORT_MAX:
            result.errors.append(RowError(number, label, f"longer than {SHORT_MAX} characters"))


def validate_learners(
    session, rows: list[dict], mapping: dict, *, school_year_id=None, adviser_user_id=None
) -> ValidationResult:
    """`adviser_user_id` scopes the Section column to that adviser's own
    sections (§3C), so an adviser can enrol their class from the file
    while a row naming someone else's section is refused. Left None — the
    Registrar and Super Admin path — any section in the year resolves."""
    result = ValidationResult()

    existing_lrns = {
        lrn for (lrn,) in session.query(Learner.lrn).filter(Learner.lrn.isnot(None)).all()
    }
    seen_in_file: dict[str, int] = {}
    # Loaded once for the whole file, not per row — a 1,200-row masterlist
    # checked one query at a time would take minutes.
    sections_by_name, ambiguous_names = _section_lookup(
        session, school_year_id, adviser_user_id
    )

    # Decided once for the whole column: 03/04/2009 is two different days,
    # and a single unambiguous value elsewhere in the file settles it for
    # every row. See detect_date_order.
    date_order = detect_date_order(row.get("birthdate") for row in rows)

    for row in rows:
        number = row.get("__row__")
        errors_before = len(result.errors)

        last_name = normalize_name(row.get("last_name"))
        first_name = normalize_name(row.get("first_name"))
        if not last_name:
            result.errors.append(RowError(number, "Last Name", "required"))
        if not first_name:
            result.errors.append(RowError(number, "First Name", "required"))

        middle_name = normalize_name(row.get("middle_name"))
        extension_name = normalize_name(row.get("extension_name"))
        guardian_name = normalize_name(row.get("guardian_name"))
        _flag_long_names(
            result, number,
            [("Last Name", last_name), ("First Name", first_name), ("Middle Name", middle_name),
             ("Extension", extension_name), ("Parent/Guardian Name", guardian_name)],
        )

        sex, sex_error = _parse_sex(row.get("sex"))
        if sex_error:
            result.errors.append(RowError(number, "Sex", sex_error))

        birthdate, date_error = parse_date(row.get("birthdate"), date_order)
        if date_error:
            result.errors.append(RowError(number, "Birthdate", date_error))
        elif birthdate is None:
            result.errors.append(RowError(number, "Birthdate", "required"))

        lrn, lrn_error = parse_lrn(row.get("lrn"))
        if lrn_error:
            result.errors.append(RowError(number, "LRN", lrn_error))
        elif lrn:
            if lrn in existing_lrns:
                result.errors.append(
                    RowError(number, "LRN", f"duplicate LRN — {lrn} already exists in the system")
                )
            elif lrn in seen_in_file:
                result.errors.append(
                    RowError(
                        number, "LRN",
                        f"duplicate LRN — same as row {seen_in_file[lrn]} in this file",
                    )
                )
            else:
                seen_in_file[lrn] = number

        # Optional. Left blank, the learner is created but not enrolled —
        # exactly what the import did before this column existed.
        section = None
        raw_section = (row.get("section") or "").strip()
        if raw_section:
            if school_year_id is None:
                result.errors.append(
                    RowError(number, "Section", "choose a school year above to enrol into a section")
                )
            elif raw_section.upper() in ambiguous_names:
                result.errors.append(
                    RowError(
                        number, "Section",
                        f"more than one section is called {raw_section!r} this school year — "
                        "rename one of them so they can be told apart",
                    )
                )
            else:
                section = sections_by_name.get(raw_section.upper())
                if section is None:
                    result.errors.append(
                        RowError(number, "Section", f"unknown section {raw_section!r}")
                    )
                elif adviser_user_id is not None and not is_advised_by(
                    section, adviser_user_id
                ):
                    # Refused per row rather than by dropping the whole
                    # column: the rest of the file still imports, and the
                    # teacher is told which rows were left out.
                    result.errors.append(
                        RowError(
                            number, "Section",
                            f"{section.name} is not one of your sections — you can only "
                            "enrol into a section you advise",
                        )
                    )
                    section = None

        # Contact cells are optional, so an unreadable one is a *warning*:
        # the learner is still created, with that field blank. Refusing the
        # row would block a valid learner over data that isn't required.
        guardian_email, email_error = clean_email(row.get("guardian_email"))
        if email_error:
            result.warnings.append(
                RowError(number, "Parent/Guardian Email", f"{email_error} — left blank")
            )
        guardian_mobile, mobile_error = normalise_mobile(row.get("guardian_mobile"))
        if mobile_error:
            result.warnings.append(
                RowError(number, "Parent/Guardian Mobile", f"{mobile_error} — left blank")
            )
        consent, consent_error = parse_consent(row.get("notices_consent"))
        if consent_error:
            result.warnings.append(
                RowError(number, "Notice Consent", f"{consent_error} — recorded as no consent")
            )
        elif consent and consent_without_contact(
            {"notices_consent": consent, "guardian_email": guardian_email,
             "guardian_mobile": guardian_mobile}
        ):
            result.warnings.append(
                RowError(number, "Notice Consent", f"{CONSENT_NEEDS_CONTACT} — recorded as no consent")
            )
            consent = False

        if len(result.errors) == errors_before:
            result.parsed.append(
                {
                    "__row__": number,
                    "last_name": last_name,
                    "first_name": first_name,
                    "middle_name": middle_name,
                    "extension_name": extension_name,
                    "sex": sex,
                    "birthdate": birthdate,
                    "lrn": lrn,
                    "guardian_name": guardian_name,
                    "guardian_email": guardian_email,
                    "guardian_mobile": guardian_mobile,
                    "notices_consent": bool(consent),
                    "section_id": section.id if section else None,
                    "grade_level_id": section.grade_level_id if section else None,
                    "school_year_id": school_year_id if section else None,
                }
            )
    return result


def commit_learners(session, parsed: list[dict], user_id=None) -> int:
    """`user_id` is stamped on each learner rather than written to the
    audit log. Attribution for a bulk add is the column, not an entry per
    row: a 1,200-learner migration would put 1,200 identical
    LEARNER_CREATED rows into a viewer that shows the most recent 200,
    burying every other kind of change behind one afternoon's import. The
    import itself is already recorded once, as DATA_IMPORTED.

    It is also what decides who may fix a typo afterwards — see
    `app.learner_access`. This parameter was accepted and discarded until
    2026-08-21.
    """
    for row in parsed:
        learner = Learner(
            last_name=row["last_name"],
            first_name=row["first_name"],
            middle_name=row["middle_name"],
            extension_name=row["extension_name"],
            sex=row["sex"],
            birthdate=row["birthdate"],
            lrn=row["lrn"],
            guardian_name=row.get("guardian_name"),
            guardian_email=row.get("guardian_email"),
            guardian_mobile=row.get("guardian_mobile"),
            notices_consent=row.get("notices_consent", False),
            # Dated when it is recorded, in school time — the host runs on
            # UTC and would date an 8 a.m. Manila import yesterday.
            notices_consent_date=(
                datetime.now(SCHOOL_TZ).date() if row.get("notices_consent") else None
            ),
            created_by_user_id=user_id,
        )
        session.add(learner)

        if not row.get("section_id"):
            continue
        # This codebase declares no ORM relationship() between models, so
        # SQLAlchemy cannot infer that the learner must be inserted before
        # the enrollment that references it by a bare FK column. Flushing
        # here makes the generated id available and fixes the insert order.
        session.flush()
        session.add(
            Enrollment(
                learner_id=learner.id,
                school_year_id=row["school_year_id"],
                grade_level_id=row["grade_level_id"],
                section_id=row["section_id"],
                enrollment_status=EnrollmentStatus.ENROLLED,
            )
        )
    return len(parsed)


# --- Term grades -----------------------------------------------------------

TERM_GRADE_COLUMNS = [
    ColumnSpec("lrn", "LRN", True, ("learnerreferencenumber",)),
    ColumnSpec("section", "Section", True, ("sectionname",)),
    ColumnSpec("subject", "Subject", True, ("subjectname", "learningarea")),
    ColumnSpec("term", "Term", True, ("termnumber", "quarter")),
    ColumnSpec("grade", "Grade", True, ("finalgrade", "rating", "termgrade")),
]


def _parse_term_number(raw: str):
    value = str(raw or "").strip().upper()
    for prefix in ("TERM", "T", "QUARTER", "Q"):
        if value.startswith(prefix):
            value = value[len(prefix):].strip()
            break
    if value in ("1", "2", "3"):
        return int(value), None
    return None, f"{raw!r} is not term 1, 2 or 3"


def validate_term_grades(
    session, rows: list[dict], mapping: dict, *, school_year_id=None
) -> ValidationResult:
    """Resolves each row against the section's actual offerings.

    The last check is the one that matters most: a grade for a subject
    the section doesn't offer *in that term* is rejected rather than
    silently written, because `section_subject_offerings` is the single
    source of truth for what a learner is graded on (rule 5).

    **Everything is scoped to one school year.** A section name is only
    unique per grade level per year, so a lookup keyed on the name alone
    matches across every year the school has ever run — and the last one
    loaded silently wins. Writing a Term 1 grade into last year's section
    of the same name produces a plausible-looking row that no report will
    ever show, which is the worst kind of wrong. Scoping is also what
    keeps this affordable: the unscoped version read every enrollment and
    every term grade in the database, which at 1,200 learners is tens of
    thousands of rows per file.
    """
    result = ValidationResult()
    if school_year_id is None:
        result.errors.append(RowError(None, None, "Choose the school year this file belongs to."))
        return result

    learners_by_lrn = {
        lrn: learner_id
        for learner_id, lrn in session.query(Learner.id, Learner.lrn)
        .filter(Learner.lrn.isnot(None))
        .all()
    }
    sections_by_name, ambiguous_names = _section_lookup(session, school_year_id)
    subjects = {s.official_name.strip().upper(): s for s in session.query(Subject).all()}
    terms_by_number = {
        t.term_number: t
        for t in session.query(Term).filter_by(school_year_id=school_year_id).all()
    }
    # (section_id, subject_id, term_id) -> offering
    offerings = {
        (o.section_id, o.subject_id, o.term_id): o
        for o in session.query(SectionSubjectOffering)
        .filter_by(school_year_id=school_year_id)
        .all()
    }
    enrollments = {
        (e.learner_id, e.section_id): e
        for e in session.query(Enrollment).filter_by(school_year_id=school_year_id).all()
    }
    # Only the keys are needed, and only for this year's enrollments —
    # loading whole TermGrade objects for the entire school was the single
    # most expensive query in the app.
    existing: set = set()
    enrollment_ids = [e.id for e in enrollments.values()]
    if enrollment_ids:
        existing = set(
            session.query(TermGrade.enrollment_id, TermGrade.section_subject_offering_id)
            .filter(TermGrade.enrollment_id.in_(enrollment_ids))
            .all()
        )

    for row in rows:
        number = row.get("__row__")
        errors_before = len(result.errors)

        lrn, lrn_error = parse_lrn(row.get("lrn"))
        if lrn_error:
            result.errors.append(RowError(number, "LRN", lrn_error))
        elif not lrn:
            result.errors.append(RowError(number, "LRN", "required"))
        elif lrn not in learners_by_lrn:
            result.errors.append(RowError(number, "LRN", f"no learner with LRN {lrn}"))

        raw_section = str(row.get("section", "")).strip()
        section = None
        if raw_section.upper() in ambiguous_names:
            result.errors.append(
                RowError(
                    number, "Section",
                    f"more than one section is called {raw_section!r} this school year — "
                    "rename one of them so they can be told apart",
                )
            )
        else:
            section = sections_by_name.get(raw_section.upper())
            if section is None:
                result.errors.append(
                    RowError(number, "Section", f"unknown section {raw_section!r}")
                )

        subject = subjects.get(str(row.get("subject", "")).strip().upper())
        if subject is None:
            result.errors.append(
                RowError(number, "Subject", f"unknown subject {row.get('subject')!r}")
            )

        term_number, term_error = _parse_term_number(row.get("term"))
        if term_error:
            result.errors.append(RowError(number, "Term", term_error))
        elif term_number not in terms_by_number:
            result.errors.append(
                RowError(number, "Term", f"this school year has no term {term_number}")
            )

        grade, grade_error = parse_grade(row.get("grade"))
        if grade_error:
            result.errors.append(RowError(number, "Grade", f"invalid grade — {grade_error}"))
        elif grade is None:
            result.errors.append(RowError(number, "Grade", "required"))

        if len(result.errors) != errors_before:
            continue

        enrollment = enrollments.get((learners_by_lrn[lrn], section.id))
        if enrollment is None:
            result.errors.append(
                RowError(number, "Section", f"learner {lrn} is not enrolled in {section.name}")
            )
            continue

        term = terms_by_number[term_number]
        offering = offerings.get((section.id, subject.id, term.id))
        if offering is None:
            result.errors.append(
                RowError(
                    number, "Subject",
                    f"{subject.official_name} is not offered to {section.name} in term "
                    f"{term_number}",
                )
            )
            continue

        result.parsed.append(
            {
                "__row__": number,
                "enrollment_id": enrollment.id,
                "offering_id": offering.id,
                "term_id": offering.term_id,
                "grade": grade,
                "replaces_existing": (enrollment.id, offering.id) in existing,
            }
        )
    return result


def commit_term_grades(session, parsed: list[dict], user_id=None) -> int:
    """Writes as DRAFT, never straight to FINALIZED — a migrated grade
    still goes through the normal submit/verify workflow (rule 7)."""
    enrollment_ids = list({row["enrollment_id"] for row in parsed})
    existing: dict = {}
    if enrollment_ids:
        existing = {
            (g.enrollment_id, g.section_subject_offering_id): g
            for g in session.query(TermGrade)
            .filter(TermGrade.enrollment_id.in_(enrollment_ids))
            .all()
        }
    for row in parsed:
        key = (row["enrollment_id"], row["offering_id"])
        grade = existing.get(key)
        if grade is None:
            # No `encoded_by_user_id` here: that column is on
            # AttendanceRecord, not TermGrade. Passing it raised TypeError
            # on every INSERT and was silently dropped on every UPDATE, so
            # importing new term grades could not work at all. Who ran the
            # import is recorded on the `import_jobs` row and in the audit
            # log; TermGrade only tracks who submitted, verified and
            # finalized, which the workflow sets later.
            session.add(
                TermGrade(
                    enrollment_id=row["enrollment_id"],
                    section_subject_offering_id=row["offering_id"],
                    term_id=row["term_id"],
                    official_grade=row["grade"],
                    status=GradeWorkflowStatus.DRAFT,
                    source=GRADE_SOURCE_IMPORT,
                )
            )
        else:
            grade.official_grade = row["grade"]
            grade.status = GradeWorkflowStatus.DRAFT
            grade.source = GRADE_SOURCE_IMPORT
            grade.version = (grade.version or 0) + 1
    return len(parsed)


def recompute_after_term_grades(session, parsed: list[dict], progress=None) -> str:
    """Refresh the derived tables the imported grades feed.

    The Gradebook does this after every save, so an import that skipped it
    would leave Grade Summary, the term cards and SF9 showing blanks for
    learners whose grades are already in the database — a silent wrong
    state rather than an error.

    It runs **after** the import's own transaction because
    `recompute_enrollment_grades_batch` commits internally. Batched into
    one call for the whole file rather than once per learner — that used
    to cost a handful of round trips *per learner* (a file covering the
    whole school took minutes), which is why the page still advises
    importing a section at a time even though this step is now a fixed
    cost regardless of file size.
    """
    enrollment_ids = list(dict.fromkeys(row["enrollment_id"] for row in parsed))
    if enrollment_ids:
        recompute_enrollment_grades_batch(session, enrollment_ids)
        if progress:
            progress(len(enrollment_ids), len(enrollment_ids))
    return f"Averages recomputed for {len(enrollment_ids)} learner(s)."


LEARNER_IMPORT = ImportSpec(
    job_type=ImportJobType.LEARNERS,
    label="Learners",
    description=(
        "The learner masterlist. Fill in the optional Section column and each "
        "learner is enrolled at the same time; leave it blank and they are "
        "created only, to be enrolled later on the Enrollment page."
    ),
    columns=LEARNER_COLUMNS,
    validate=validate_learners,
    commit=commit_learners,
    needs_school_year=True,
    school_year_help=(
        "Only used by the optional Section column. Learners with no section are "
        "created without being enrolled."
    ),
)

# --- Parent/guardian contacts (update by LRN) -----------------------------

CONTACT_COLUMNS = [
    ColumnSpec("lrn", "LRN", True, ("learnerreferencenumber",)),
    *GUARDIAN_COLUMNS,
]


def validate_learner_contacts(
    session, rows: list[dict], mapping: dict, *, adviser_user_id=None, editable_ids=None
) -> ValidationResult:
    """Contact details for learners **already in the system**, matched by LRN.

    The learner import only ever INSERTs, so without this every existing
    learner's contact would have to be typed one at a time.

    **A blank cell leaves the stored value alone.** A file carrying only a
    mobile column must not wipe everyone's email. Clearing a value is done
    on the learner's own form.

    Unlike in the learner import, an unreadable contact cell is an *error*
    here: the contact is the whole of the row, so there is nothing else to
    save from it.

    `adviser_user_id`/`editable_ids` scope it as the Masterlist does (§3C):
    an adviser may update only the learners they may edit. None is unscoped.
    """
    result = ValidationResult()

    lrns = {lrn for lrn, _ in (parse_lrn(row.get("lrn")) for row in rows) if lrn}
    # One query for the whole file.
    learners = (
        {row.lrn: row for row in session.query(Learner).filter(Learner.lrn.in_(lrns)).all()}
        if lrns
        else {}
    )

    seen_in_file: dict[str, int] = {}
    for row in rows:
        number = row.get("__row__")
        errors_before = len(result.errors)

        learner = None
        lrn, lrn_error = parse_lrn(row.get("lrn"))
        if lrn_error:
            result.errors.append(RowError(number, "LRN", lrn_error))
        elif not lrn:
            result.errors.append(RowError(number, "LRN", "required"))
        elif lrn in seen_in_file:
            result.errors.append(
                RowError(number, "LRN", f"same LRN as row {seen_in_file[lrn]} in this file")
            )
        else:
            seen_in_file[lrn] = number
            learner = learners.get(lrn)
            if learner is None:
                result.errors.append(RowError(number, "LRN", f"no learner with LRN {lrn}"))
            elif not may_edit(learner.id, editable_ids or set(), adviser_user_id):
                result.errors.append(RowError(number, "LRN", f"{lrn} is not one of your learners"))

        email, email_error = clean_email(row.get("guardian_email"))
        if email_error:
            result.errors.append(RowError(number, "Parent/Guardian Email", email_error))
        mobile, mobile_error = normalise_mobile(row.get("guardian_mobile"))
        if mobile_error:
            result.errors.append(RowError(number, "Parent/Guardian Mobile", mobile_error))
        consent, consent_error = parse_consent(row.get("notices_consent"))
        if consent_error:
            result.errors.append(RowError(number, "Notice Consent", consent_error))
        _flag_long_names(
            result, number, [("Parent/Guardian Name", normalize_name(row.get("guardian_name")))]
        )

        if len(result.errors) != errors_before:
            continue
        # Only what the row actually states; a blank cell is left out, so
        # it can't overwrite anything.
        changes = {
            "guardian_name": normalize_name(row.get("guardian_name")),
            "guardian_email": email,
            "guardian_mobile": mobile,
            "notices_consent": consent,
        }
        changes = {k: v for k, v in changes.items() if v is not None}
        # Against what the save would leave, stored values included: "Yes"
        # for a learner with no email or mobile on file *or* in the file.
        if consent_without_contact({**contact_values(learner), **changes}):
            result.errors.append(RowError(number, "Notice Consent", CONSENT_NEEDS_CONTACT))
            continue
        result.parsed.append(
            {
                "__row__": number,
                "__changes__": changes,
                "learner_id": learner.id,
                "lrn": lrn,
                "name": f"{learner.last_name}, {learner.first_name}",
                **changes,
            }
        )
    return result


def commit_learner_contacts(session, parsed: list[dict], user_id=None) -> int:
    """Returns the number of learners whose contact actually changed.

    Audited **per learner**, with old and new values, unlike the learner
    import's single DATA_IMPORTED entry (see `commit_learners`). That one
    creates rows, so there is no old value to lose; this overwrites where a
    term card will be emailed, and the previous address is exactly what
    someone needs if a card goes to the wrong parent. Only learners whose
    values differ are logged, so re-uploading the same file adds nothing.
    """
    ids = [row["learner_id"] for row in parsed]
    if not ids:
        return 0
    learners = {
        learner.id: learner
        for learner in session.query(Learner).filter(Learner.id.in_(ids)).all()
    }
    # In school time; the host runs on UTC.
    today = datetime.now(SCHOOL_TZ).date()
    changed = 0
    for row in parsed:
        learner = learners.get(row["learner_id"])
        if learner is None:
            continue
        previous = contact_values(learner)
        apply_contact(learner, row["__changes__"], today)
        was, now = audit_service.changes(previous, contact_values(learner))
        if was:
            audit_service.record(
                session,
                action=audit_service.LEARNER_CONTACT_CHANGED,
                object_type="learners",
                object_id=learner.id,
                user_id=user_id,
                previous=was,
                new=now,
            )
            changed += 1
    return changed


# --- SHS admission records (update by LRN) --------------------------------

# The same fields, in the same order, as the Learner Masterlist's
# "Save admission record" form (`_admission_record_form`), plus the
# completion date the record also holds. (field, header, kind, limit).
ADMISSION_FIELDS = [
    ("date_of_shs_admission", "Date of SHS Admission", "date", None),
    ("junior_high_school_completer", "JHS Completer", "bool", None),
    ("junior_high_school_general_average", "JHS General Average", "number", None),
    ("high_school_completer", "HS Completer", "bool", None),
    ("high_school_general_average", "HS General Average", "number", None),
    ("high_school_completion_date", "HS Completion Date", "date", None),
    ("previous_school_name", "Previous School Name", "text", ADDRESS_MAX),
    ("previous_school_address", "Previous School Address", "text", ADDRESS_MAX),
    ("pept_passer", "PEPT Passer", "bool", None),
    ("pept_rating", "PEPT Rating", "number", None),
    ("pept_examination_date", "PEPT Exam Date", "date", None),
    ("als_ae_passer", "ALS A&E Passer", "bool", None),
    ("als_ae_rating", "ALS A&E Rating", "number", None),
    ("als_ae_examination_date", "ALS A&E Exam Date", "date", None),
    ("clc_name", "CLC Name", "text", ADDRESS_MAX),
    ("clc_address", "CLC Address", "text", ADDRESS_MAX),
    ("other_eligibility_notes", "Other Eligibility Notes", "text", NOTE_MAX),
]

_ADMISSION_ALIASES = {
    "date_of_shs_admission": ("dateofadmission", "shsadmissiondate", "admissiondate"),
    "junior_high_school_completer": ("jhscompleter", "juniorhighschoolcompleter"),
    "junior_high_school_general_average": ("jhsgeneralaverage", "jhsga", "jhsaverage"),
    "high_school_completer": ("hscompleter", "highschoolcompleter"),
    "high_school_general_average": ("hsgeneralaverage", "hsga", "hsaverage"),
    "high_school_completion_date": ("hscompletiondate", "dateofhscompletion"),
    "previous_school_name": ("previousschool", "lastschoolattended"),
    "pept_examination_date": ("peptexaminationdate", "peptdate"),
    "als_ae_examination_date": ("alsaeexaminationdate", "alsdate", "alsaedate"),
}

ADMISSION_COLUMNS = [
    ColumnSpec("lrn", "LRN", True, ("learnerreferencenumber",)),
    *(
        ColumnSpec(field_name, label, False, _ADMISSION_ALIASES.get(field_name, ()))
        for field_name, label, _kind, _limit in ADMISSION_FIELDS
    ),
]

_YES = {"y", "yes", "true", "1", "oo"}
_NO = {"n", "no", "false", "0", "hindi"}
GRADE_12_MESSAGE = "SHS admission is for Grade 11 learners"


def _parse_yes_no(raw) -> tuple[bool | None, str | None]:
    value = str(raw if raw is not None else "").strip().lower()
    if not value:
        return None, None
    if value in _YES:
        return True, None
    if value in _NO:
        return False, None
    return None, f"{str(raw).strip()!r} is not Yes or No"


def _parse_percentage(raw) -> tuple[Decimal | None, str | None]:
    """A rating or general average: a number from 0 to 100, kept to the two
    decimals the column holds. Blank is not an error and not a zero."""
    text = str(raw if raw is not None else "").strip().replace(",", ".")
    if not text:
        return None, None
    try:
        value = Decimal(text)
    except InvalidOperation:
        return None, f"{str(raw).strip()!r} is not a number"
    if not value.is_finite() or not (0 <= value <= 100):
        return None, f"{str(raw).strip()} is outside 0-100"
    return value.quantize(Decimal("0.01")), None


def validate_learner_admissions(
    session, rows: list[dict], mapping: dict, *, adviser_user_id=None, editable_ids=None
) -> ValidationResult:
    """SHS admission records for learners **already in the system**,
    matched by LRN. Same fields and meaning as the Masterlist form.

    Two queries for the lookups however long the file is (learners, then
    the Grade 12 set). **A blank cell leaves the stored value alone** and
    a non-blank one overwrites, so a field can only be cleared on the
    form. A learner with no record gets one at commit.

    Refused per row, never skipped silently: an unknown LRN, a learner the
    viewer may not edit (`editable_ids`; no adviser means unscoped), and a
    learner whose newest enrollment is Grade 12 — the form is not offered
    for them either, and both use `learner_access.grade_12_learner_ids`.
    """
    result = ValidationResult()

    lrns = {lrn for lrn, _ in (parse_lrn(row.get("lrn")) for row in rows) if lrn}
    learners = (
        {row.lrn: row for row in session.query(Learner).filter(Learner.lrn.in_(lrns)).all()}
        if lrns
        else {}
    )
    grade_12 = grade_12_learner_ids(session, [learner.id for learner in learners.values()])

    # One reading order for every slash date in the file, decided from all
    # the date columns together (same reasoning as the learner import).
    date_fields = [name for name, _label, kind, _limit in ADMISSION_FIELDS if kind == "date"]
    order = detect_date_order(row.get(name) for row in rows for name in date_fields)

    seen_in_file: dict[str, int] = {}
    for row in rows:
        number = row.get("__row__")
        errors_before = len(result.errors)

        learner = None
        lrn, lrn_error = parse_lrn(row.get("lrn"))
        if lrn_error:
            result.errors.append(RowError(number, "LRN", lrn_error))
        elif not lrn:
            result.errors.append(RowError(number, "LRN", "required"))
        elif lrn in seen_in_file:
            result.errors.append(
                RowError(number, "LRN", f"same LRN as row {seen_in_file[lrn]} in this file")
            )
        else:
            seen_in_file[lrn] = number
            learner = learners.get(lrn)
            if learner is None:
                result.errors.append(RowError(number, "LRN", f"no learner with LRN {lrn}"))
            elif not may_edit(learner.id, editable_ids or set(), adviser_user_id):
                result.errors.append(RowError(number, "LRN", f"{lrn} is not one of your learners"))
            elif learner.id in grade_12:
                result.errors.append(RowError(number, "LRN", GRADE_12_MESSAGE))

        changes: dict = {}
        for name, label, kind, limit in ADMISSION_FIELDS:
            raw = row.get(name)
            if kind == "bool":
                value, problem = _parse_yes_no(raw)
            elif kind == "number":
                value, problem = _parse_percentage(raw)
            elif kind == "date":
                value, problem = parse_date(raw, order)
            else:
                value = str(raw if raw is not None else "").strip() or None
                problem = (
                    f"longer than {limit} characters" if value and len(value) > limit else None
                )
            if problem:
                result.errors.append(RowError(number, label, problem))
            elif value is not None:
                changes[name] = value

        if len(result.errors) != errors_before:
            continue
        if not changes:
            result.warnings.append(
                RowError(number, None, "no admission details in this row, so nothing to update")
            )
            continue
        result.parsed.append(
            {
                "__row__": number,
                "__changes__": changes,
                "learner_id": learner.id,
                "lrn": lrn,
                "name": f"{learner.last_name}, {learner.first_name}",
                **changes,
            }
        )
    return result


def commit_learner_admissions(session, parsed: list[dict], user_id=None) -> int:
    """Returns the number of learners whose admission record changed.

    One query loads every existing record; a learner with none gets a new
    row. Audited per learner as LEARNER_ADMISSION_CHANGED with the old and
    new value of each field that differs, so re-uploading the same file
    adds nothing. The form's entry is a thin summary; this one overwrites
    many learners at once, so the previous values are worth keeping.
    """
    ids = [row["learner_id"] for row in parsed]
    if not ids:
        return 0
    records = {
        record.learner_id: record
        for record in session.query(LearnerAdmissionRecord)
        .filter(LearnerAdmissionRecord.learner_id.in_(ids))
        .all()
    }
    changed = 0
    for row in parsed:
        record = records.get(row["learner_id"])
        created = record is None
        if created:
            record = LearnerAdmissionRecord(learner_id=row["learner_id"])
            session.add(record)
            records[row["learner_id"]] = record
        fields = row["__changes__"]
        previous = {name: getattr(record, name, None) for name in fields}
        for name, value in fields.items():
            setattr(record, name, value)
        was, now = audit_service.changes(previous, fields)
        if was or created:
            audit_service.record(
                session,
                action=audit_service.LEARNER_ADMISSION_CHANGED,
                object_type="learners",
                object_id=row["learner_id"],
                user_id=user_id,
                previous=was,
                new={"created": created, **now},
            )
            changed += 1
    return changed


def admission_template_bytes() -> bytes:
    """A blank .xlsx: the headers, plus one example row.

    The example's LRN is all zeros, which no learner has, so a template
    uploaded with the example still in it fails that row ("no learner with
    LRN") instead of writing anything. LRN is a text cell, so Excel keeps
    the zeros.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "SHS admission"
    sheet.append([column.label for column in ADMISSION_COLUMNS])
    example = {
        "lrn": "000000000000",
        "date_of_shs_admission": date(2026, 6, 8),
        "junior_high_school_completer": "Yes",
        "junior_high_school_general_average": 88.5,
        "high_school_completer": "No",
        "previous_school_name": "Example National High School",
        "previous_school_address": "Example Street, Example City",
        "pept_passer": "No",
        "als_ae_passer": "No",
        "other_eligibility_notes": "Delete this example row before uploading.",
    }
    sheet.append([example.get(column.field) for column in ADMISSION_COLUMNS])
    kinds = {name: kind for name, _label, kind, _limit in ADMISSION_FIELDS}
    for index, column in enumerate(ADMISSION_COLUMNS, start=1):
        letter = openpyxl.utils.get_column_letter(index)
        sheet.cell(row=1, column=index).font = openpyxl.styles.Font(bold=True)
        sheet.column_dimensions[letter].width = max(16, len(column.label) + 4)
        if column.field == "lrn":
            sheet.cell(row=2, column=index).number_format = "@"
        elif kinds.get(column.field) == "date":
            sheet.cell(row=2, column=index).number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# Example rows are offered only where an example left in the file cannot
# write anything: an all-zero LRN matches no learner, so the row fails
# validation instead. The learner and term-grade imports have no such safe
# example (a filled-in learner row would be created), so they get headers only.
_TEMPLATE_EXAMPLES = {
    ImportJobType.LEARNER_CONTACTS: {
        "lrn": "000000000000",
        "guardian_name": "Delete this example row before uploading.",
        "guardian_email": "parent@example.com",
        "guardian_mobile": "09171234567",
        "notices_consent": "Yes",
    },
}
_TEXT_FIELDS = {"lrn", "guardian_mobile"}


def spec_template_bytes(spec) -> bytes:
    """A blank .xlsx for any import spec: headers, plus a Columns sheet.

    Built from `spec.columns`, so the template cannot drift from what the
    importer actually reads. LRN and mobile columns are formatted as text so
    Excel keeps leading zeros.
    """
    if spec.job_type == ImportJobType.LEARNER_ADMISSIONS:
        return admission_template_bytes()  # its own, with typed date cells
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append([column.label for column in spec.columns])
    example = _TEMPLATE_EXAMPLES.get(spec.job_type)
    if example:
        sheet.append([example.get(column.field) for column in spec.columns])
    bold = openpyxl.styles.Font(bold=True)
    for index, column in enumerate(spec.columns, start=1):
        letter = openpyxl.utils.get_column_letter(index)
        sheet.cell(row=1, column=index).font = bold
        sheet.column_dimensions[letter].width = max(16, len(column.label) + 4)
        if column.field in _TEXT_FIELDS:
            # Pre-format enough rows that typed LRNs keep their zeros.
            for row in range(2, 502):
                sheet.cell(row=row, column=index).number_format = "@"
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet("Columns")
    notes.append(["Column", "Required"])
    for cell in notes[1]:
        cell.font = bold
    for column in spec.columns:
        notes.append([column.label, "yes" if column.required else "optional"])
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 12
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


LEARNER_ADMISSION_IMPORT = ImportSpec(
    job_type=ImportJobType.LEARNER_ADMISSIONS,
    label="SHS admission records (update existing learners)",
    description=(
        "Adds or updates the SHS admission record of Grade 11 learners already in "
        "the system, matched by LRN. A blank cell leaves what is stored alone. "
        "Grade 12 learners are refused. Nobody is created or enrolled."
    ),
    columns=ADMISSION_COLUMNS,
    validate=validate_learner_admissions,
    commit=commit_learner_admissions,
)


TERM_GRADE_IMPORT = ImportSpec(
    job_type=ImportJobType.TERM_GRADES,
    label="Term grades",
    description=(
        "One row per learner, subject and term. Grades land as DRAFT and still "
        "go through the normal submit/verify workflow. Re-importing the same "
        "learner/subject/term updates the existing grade rather than duplicating it. "
        "Import one section at a time — averages are recalculated for every learner "
        "in the file afterwards, and that is the slow part."
    ),
    columns=TERM_GRADE_COLUMNS,
    validate=validate_term_grades,
    commit=commit_term_grades,
    needs_school_year=True,
    school_year_help=(
        "Sections, subjects and terms are all matched inside this year. A section "
        "name on its own can exist in more than one year, so this is what stops a "
        "grade landing on the wrong one."
    ),
    after_commit=recompute_after_term_grades,
)

LEARNER_CONTACT_IMPORT = ImportSpec(
    job_type=ImportJobType.LEARNER_CONTACTS,
    label="Parent/guardian contacts (update existing learners)",
    description=(
        "Adds or updates the parent/guardian name, email, mobile and notice consent "
        "of learners already in the system, matched by LRN. A blank cell leaves "
        "what is stored alone. Nobody is created or enrolled."
    ),
    columns=CONTACT_COLUMNS,
    validate=validate_learner_contacts,
    commit=commit_learner_contacts,
)

SPECS = {
    spec.job_type: spec
    for spec in (
        LEARNER_IMPORT,
        LEARNER_CONTACT_IMPORT,
        LEARNER_ADMISSION_IMPORT,
        TERM_GRADE_IMPORT,
    )
}
